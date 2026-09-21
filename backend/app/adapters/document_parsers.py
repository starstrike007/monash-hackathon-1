from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from app.adapters.dataset_loader import DatasetLoader
from app.api.schemas.common import DocumentType


@dataclass
class ParsedDocument:
    path: str
    text: str = ""
    table_rows: list[list[str]] = field(default_factory=list)
    # Locations are aligned with text lines where possible. Table locations
    # also retain sheet/page/table and cell coordinates for review evidence.
    locations: list[dict[str, Any]] = field(default_factory=list)
    readable: bool = True
    error: str | None = None
    error_code: str | None = None
    missing: bool = False
    document_type: DocumentType = DocumentType.UNKNOWN


def detect_document_type(text: str) -> DocumentType:
    upper = text.upper()
    if re.search(
        r"BILL\s+OF\s+LADING\s+INSTRUCTION|\bBL\s+INSTRUCTION\b|\bSI\s+INSTRUCTION\b|SHIPPING\s+INSTRUCTION",
        upper,
    ):
        return DocumentType.SHIPPING_INSTRUCTION
    if re.search(r"BILL\s+OF\s+LADING|\bB/L\b|DRAFT\s+BL", upper):
        return DocumentType.BILL_OF_LADING
    if re.search(r"SHIPPER/EXPORTER|NOTIFY\s+PARTY", upper):
        return DocumentType.SHIPPING_INSTRUCTION
    if "COMMERCIAL INVOICE" in upper or re.search(r"\bINVOICE\b", upper):
        return DocumentType.COMMERCIAL_INVOICE
    if "PACKING LIST" in upper:
        return DocumentType.PACKING_LIST
    if "CERTIFICATE OF ORIGIN" in upper:
        return DocumentType.CERTIFICATE_OF_ORIGIN
    return DocumentType.UNKNOWN


def _clean_cell(value: Any) -> str:
    if value is None:
        return ""
    # Keep paragraph breaks inside a Word/Excel cell.  They are meaningful
    # structure: a party name followed by its address must not become one
    # extracted legal-name value merely because the cell contains newlines.
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    lines = [
        re.sub(r"[ \t]+", " ", line).strip()
        for line in text.split("\n")
    ]
    return "\n".join(line for line in lines if line)


def _looks_garbled_text(text: str) -> bool:
    """Return true for PDF text that cannot be trusted as source evidence."""

    non_space = [character for character in text if not character.isspace()]
    if not non_space:
        return True
    if "\ufffd" in text or "(cid:" in text.lower():
        return True
    controls = sum(
        1
        for character in non_space
        if ord(character) < 32 and character not in {"\t", "\n", "\r"}
    )
    if controls / len(non_space) > 0.01:
        return True
    # A text layer containing only drawing/punctuation artifacts is not useful
    # evidence even if pdfplumber returned a non-empty string.
    return not any(character.isalnum() for character in non_space)


def _parse_pdf(raw: bytes) -> tuple[str, list[list[str]], list[dict[str, Any]]]:
    try:
        import pdfplumber
    except ImportError as exc:
        raise RuntimeError("pdfplumber is not installed") from exc

    text_lines: list[str] = []
    table_rows: list[list[str]] = []
    locations: list[dict[str, Any]] = []
    with pdfplumber.open(io.BytesIO(raw)) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            page_text = page.extract_text() or ""
            for line_number, line in enumerate(page_text.splitlines(), start=1):
                text_lines.append(line)
                locations.append({"page": page_number, "line": line_number})

            # Keep tables as structured rows in addition to the readable text.
            # A missing text layer still remains unreadable by policy below.
            for table_number, table in enumerate(page.extract_tables() or [], start=1):
                for row_number, row in enumerate(table, start=1):
                    cells = [_clean_cell(cell) for cell in (row or [])]
                    if not any(cells):
                        continue
                    table_rows.append(cells)
                    locations.append(
                        {
                            "page": page_number,
                            "table": table_number,
                            "row": row_number,
                            "cells": [
                                {"cell": column_number, "value": value}
                                for column_number, value in enumerate(cells, start=1)
                            ],
                        }
                    )
    return "\n".join(text_lines), table_rows, locations


def _parse_docx(raw: bytes) -> tuple[str, list[list[str]], list[dict[str, Any]]]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("python-docx is not installed") from exc

    document = Document(io.BytesIO(raw))
    text_lines: list[str] = []
    table_rows: list[list[str]] = []
    locations: list[dict[str, Any]] = []
    for paragraph_number, paragraph in enumerate(document.paragraphs, start=1):
        value = _clean_cell(paragraph.text)
        if not value:
            continue
        text_lines.append(value)
        locations.append({"paragraph": paragraph_number})

    for table_number, table in enumerate(document.tables, start=1):
        for row_number, row in enumerate(table.rows, start=1):
            cells = [_clean_cell(cell.text) for cell in row.cells]
            if not any(cells):
                continue
            table_rows.append(cells)
            text_lines.append(" | ".join(cells))
            locations.append(
                {
                    "table": table_number,
                    "row": row_number,
                    "cells": [
                        {"cell": column_number, "value": value}
                        for column_number, value in enumerate(cells, start=1)
                    ],
                }
            )
    return "\n".join(text_lines), table_rows, locations


def _parse_xlsx(raw: bytes) -> tuple[str, list[list[str]], list[dict[str, Any]]]:
    try:
        import openpyxl
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("openpyxl is not installed") from exc

    workbook = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    text_lines: list[str] = []
    table_rows: list[list[str]] = []
    locations: list[dict[str, Any]] = []
    try:
        for sheet in workbook.worksheets:
            for row_number, row in enumerate(sheet.iter_rows(), start=1):
                values = [_clean_cell(cell.value) for cell in row]
                if not any(values):
                    continue
                table_rows.append(values)
                text_lines.append(" | ".join(values))
                locations.append(
                    {
                        "sheet": sheet.title,
                        "row": row_number,
                        "cells": [
                            {
                                "cell": f"{get_column_letter(column_number)}{row_number}",
                                "value": value,
                            }
                            for column_number, value in enumerate(values, start=1)
                            if value
                        ],
                    }
                )
    finally:
        workbook.close()
    return "\n".join(text_lines), table_rows, locations


def _parse_txt(raw: bytes) -> tuple[str, list[list[str]], list[dict[str, Any]]]:
    text = raw.decode("utf-8", errors="replace")
    lines = text.splitlines()
    table_rows = [[line] for line in lines if line.strip()]
    locations = [{"line": line_number} for line_number, _ in enumerate(lines, start=1)]
    return text, table_rows, locations


def parse_attachment(loader: DatasetLoader, relative_path: str) -> ParsedDocument:
    extension = Path(relative_path).suffix.lower()
    try:
        raw = loader.read_attachment_bytes(relative_path)
        if extension == ".txt":
            text, table_rows, locations = _parse_txt(raw)
            parsed = ParsedDocument(
                path=relative_path,
                text=text,
                table_rows=table_rows,
                locations=locations,
                readable=bool(text.strip()),
            )
        elif extension == ".pdf":
            text, table_rows, locations = _parse_pdf(raw)
            garbled = _looks_garbled_text(text)
            parsed = ParsedDocument(
                path=relative_path,
                text=text,
                table_rows=table_rows,
                locations=locations,
                readable=bool(text.strip()) and not garbled,
                error_code="unreadable_pdf" if not text.strip() or garbled else None,
            )
        elif extension == ".docx":
            text, table_rows, locations = _parse_docx(raw)
            parsed = ParsedDocument(
                path=relative_path,
                text=text,
                table_rows=table_rows,
                locations=locations,
                readable=bool(text.strip()),
                error_code="unreadable_document" if not text.strip() else None,
            )
        elif extension == ".xlsx":
            text, table_rows, locations = _parse_xlsx(raw)
            parsed = ParsedDocument(
                path=relative_path,
                text=text,
                table_rows=table_rows,
                locations=locations,
                readable=bool(text.strip()),
                error_code="unreadable_document" if not text.strip() else None,
            )
        else:
            parsed = ParsedDocument(
                path=relative_path,
                readable=False,
                error="Unsupported attachment extension",
                error_code="unsupported_extension",
            )
    except FileNotFoundError:
        parsed = ParsedDocument(
            path=relative_path,
            readable=False,
            missing=True,
            error="Attachment not found",
            error_code="missing_attachment",
        )
    except Exception as exc:  # parser failures are expected operational errors
        parsed = ParsedDocument(
            path=relative_path,
            readable=False,
            error=type(exc).__name__,
            error_code="parse_error",
        )

    if not parsed.readable:
        parsed.document_type = DocumentType.UNREADABLE
    else:
        parsed.document_type = detect_document_type(parsed.text)
    return parsed
