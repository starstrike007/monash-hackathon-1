from __future__ import annotations

import io
from pathlib import Path
from typing import Any

from app.adapters.dataset_loader import DatasetLoader


def _view_txt(raw: bytes) -> dict[str, Any]:
    text = raw.decode("utf-8", errors="replace")
    lines = text.replace("\r", "").split("\n")
    return {"type": "txt", "lines": lines}


def _view_pdf(raw: bytes, route_path: str) -> dict[str, Any]:
    import pdfplumber

    with pdfplumber.open(io.BytesIO(raw)) as pdf:
        page_count = len(pdf.pages)
    return {
        "type": "pdf",
        "page_count": page_count,
        "pages": [
            {"page": index, "image_url": f"/api/attachments/{route_path}/pages/{index}"}
            for index in range(1, page_count + 1)
        ],
    }


def _view_xlsx(raw: bytes) -> dict[str, Any]:
    import openpyxl
    from openpyxl.utils import get_column_letter

    workbook = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    sheets = []
    for sheet in workbook.worksheets:
        max_row = sheet.max_row or 0
        max_col = sheet.max_column or 0
        rows: list[list[str]] = []
        for row in sheet.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
            rows.append(["" if cell.value is None else str(cell.value) for cell in row])
        sheets.append(
            {
                "name": sheet.title,
                "column_letters": [get_column_letter(i) for i in range(1, max_col + 1)],
                "rows": rows,
            }
        )
    return {"type": "xlsx", "sheets": sheets}


def _view_docx(raw: bytes) -> dict[str, Any]:
    from docx import Document

    document = Document(io.BytesIO(raw))
    paragraphs = [
        {"paragraph_index": index, "text": paragraph.text}
        for index, paragraph in enumerate(document.paragraphs)
        if paragraph.text.strip()
    ]
    tables = [
        [[cell.text.strip() for cell in row.cells] for row in table.rows] for table in document.tables
    ]
    return {"type": "docx", "note": "Extracted view", "paragraphs": paragraphs, "tables": tables}


def build_document_view(loader: DatasetLoader, relative_path: str, route_path: str) -> dict[str, Any]:
    """`relative_path` is the loader-relative path (e.g. "attachments/foo.pdf");
    `route_path` is the bare path used in the public /api/attachments/{route_path}
    URLs (no "attachments/" prefix), matching the existing download endpoint."""

    extension = Path(relative_path).suffix.lower()
    raw = loader.read_attachment_bytes(relative_path)
    if extension == ".txt":
        return _view_txt(raw)
    if extension == ".pdf":
        return _view_pdf(raw, route_path)
    if extension == ".xlsx":
        return _view_xlsx(raw)
    if extension == ".docx":
        return _view_docx(raw)
    raise ValueError(f"Unsupported extension for viewer: {extension}")
