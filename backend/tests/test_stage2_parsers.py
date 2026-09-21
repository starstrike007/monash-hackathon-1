from __future__ import annotations

from pathlib import Path

from docx import Document
from openpyxl import Workbook

from app.adapters.dataset_loader import DatasetLoader
from app.adapters.document_parsers import parse_attachment
from app.api.schemas.common import CanonicalField, DocumentRole, DocumentType, ExtractionState
from app.pipeline.extract import extract_document


def _make_pdf(lines: list[str]) -> bytes:
    stream_lines = ["BT", "/F1 10 Tf", "72 760 Td"]
    for line in lines:
        escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream_lines.append(f"({escaped}) Tj")
        stream_lines.append("0 -14 Td")
    stream_lines.append("ET")
    stream = "\n".join(stream_lines).encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    payload = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(payload))
        payload.extend(f"{number} 0 obj\n".encode("ascii"))
        payload.extend(obj)
        payload.extend(b"\nendobj\n")
    xref_offset = len(payload)
    payload.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    payload.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        payload.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    payload.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode("ascii")
    )
    return bytes(payload)


def _loader(tmp_path: Path) -> DatasetLoader:
    data_dir = tmp_path / "data"
    (data_dir / "attachments").mkdir(parents=True)
    return DatasetLoader(data_dir)


def test_synthetic_txt_docx_xlsx_and_pdf_share_structured_representation(tmp_path):
    loader = _loader(tmp_path)
    attachment_dir = loader.attachments_dir

    (attachment_dir / "sample.txt").write_text(
        "SHIPPING INSTRUCTION\nShipper: Meridian Pulp\nPort of Loading: Port Klang (MYPKG)\n",
        encoding="utf-8",
    )

    docx = Document()
    docx.add_paragraph("BILL OF LADING")
    table = docx.add_table(rows=0, cols=2)
    for label, value in (("发货人", "Meridian Pulp"), ("箱数", "3 x 40HC"), ("毛重", "22000 KG")):
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
    docx.save(attachment_dir / "sample.docx")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SI"
    sheet.append(["SHIPPING INSTRUCTION", None])
    sheet.append(["Notify Party", "Yangtze Logistics"])
    sheet.append(["Port of Discharge", "Shanghai (CNSHA)"])
    notes = workbook.create_sheet("Notes")
    notes["B2"] = "supporting sheet"
    workbook.save(attachment_dir / "sample.xlsx")

    (attachment_dir / "sample.pdf").write_bytes(
        _make_pdf(["BILL OF LADING", "Shipper: Meridian Pulp", "Container Count: 3 x 40HC"])
    )

    parsed = {
        extension: parse_attachment(loader, f"attachments/sample{extension}")
        for extension in (".txt", ".docx", ".xlsx", ".pdf")
    }
    assert all(document.readable for document in parsed.values())
    assert all(document.text.strip() for document in parsed.values())
    assert all(isinstance(document.table_rows, list) for document in parsed.values())
    assert all(parsed[extension].table_rows for extension in (".txt", ".docx", ".xlsx"))
    assert all(document.locations for document in parsed.values())
    assert parsed[".xlsx"].locations[0]["sheet"] == "SI"
    assert any(cell["cell"] == "B2" for location in parsed[".xlsx"].locations for cell in location["cells"])
    assert any(location.get("page") == 1 for location in parsed[".pdf"].locations)

    docx_extraction = extract_document(parsed[".docx"], DocumentRole.BL)
    assert next(field for field in docx_extraction.fields if field.field_name == CanonicalField.SHIPPER).state == ExtractionState.FOUND
    assert next(field for field in docx_extraction.fields if field.field_name == CanonicalField.CONTAINER_COUNT).normalized_value == "3"
    assert next(field for field in docx_extraction.fields if field.field_name == CanonicalField.GROSS_WEIGHT_KG).normalized_value == "22000"

    xlsx_extraction = extract_document(parsed[".xlsx"], DocumentRole.SI)
    notify = next(field for field in xlsx_extraction.fields if field.field_name == CanonicalField.NOTIFY_PARTY)
    port = next(field for field in xlsx_extraction.fields if field.field_name == CanonicalField.PORT_OF_DISCHARGE)
    assert notify.state == ExtractionState.FOUND
    assert notify.evidence is not None and notify.evidence.sheet == "SI"
    assert port.normalized_value == "CNSHA"


def test_synthetic_pdf_without_text_layer_is_unreadable_and_not_unknown(tmp_path):
    loader = _loader(tmp_path)
    (loader.attachments_dir / "empty.pdf").write_bytes(_make_pdf([]))

    parsed = parse_attachment(loader, "attachments/empty.pdf")

    assert parsed.readable is False
    assert parsed.document_type == DocumentType.UNREADABLE
    assert parsed.error_code == "unreadable_pdf"


def test_synthetic_pdf_with_cid_garbling_is_unreadable(tmp_path):
    loader = _loader(tmp_path)
    (loader.attachments_dir / "garbled.pdf").write_bytes(_make_pdf(["(cid:123)"]))

    parsed = parse_attachment(loader, "attachments/garbled.pdf")

    assert parsed.readable is False
    assert parsed.document_type == DocumentType.UNREADABLE
    assert parsed.error_code == "unreadable_pdf"


def test_multiline_docx_value_cell_keeps_name_separate_from_address(tmp_path):
    loader = _loader(tmp_path)
    document = Document()
    document.add_paragraph("BILL OF LADING")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Shipper"
    table.cell(0, 1).text = "Meridian Pulp\n77 Harbour Road\nSingapore"
    document.save(loader.attachments_dir / "multiline.docx")

    parsed = parse_attachment(loader, "attachments/multiline.docx")
    extraction = extract_document(parsed, DocumentRole.BL)
    shipper = next(field for field in extraction.fields if field.field_name == CanonicalField.SHIPPER)

    assert shipper.state == ExtractionState.FOUND
    assert shipper.raw_value == "Meridian Pulp"
    assert "77 Harbour Road" not in shipper.raw_value


def test_colon_in_later_table_value_does_not_override_first_value_cell(tmp_path):
    loader = _loader(tmp_path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["SHIPPING INSTRUCTION", None])
    sheet.append(["Shipper", "Meridian Pulp | P.O. BOX: 123"])
    workbook.save(loader.attachments_dir / "value-colon.xlsx")

    parsed = parse_attachment(loader, "attachments/value-colon.xlsx")
    extraction = extract_document(parsed, DocumentRole.SI)
    shipper = next(field for field in extraction.fields if field.field_name == CanonicalField.SHIPPER)

    assert shipper.state == ExtractionState.FOUND
    assert shipper.raw_value == "Meridian Pulp"


def test_pdf_table_header_does_not_become_gross_weight_value(tmp_path):
    loader = _loader(tmp_path)
    (loader.attachments_dir / "header.pdf").write_bytes(
        _make_pdf(
            [
                "BILL OF LADING",
                "CONTAINER NO. DESCRIPTION GROSS WEIGHT (KG)",
                "ABC123 40HC PAPER 23,702",
                "TOTAL GROSS WEIGHT: 23,702 KG",
            ]
        )
    )

    parsed = parse_attachment(loader, "attachments/header.pdf")
    extraction = extract_document(parsed, DocumentRole.BL)
    gross_weight = next(
        field for field in extraction.fields if field.field_name == CanonicalField.GROSS_WEIGHT_KG
    )

    assert gross_weight.state == ExtractionState.FOUND
    assert gross_weight.normalized_value == "23702"
    assert gross_weight.evidence is not None
    assert "TOTAL GROSS WEIGHT" in gross_weight.evidence.snippet


def test_missing_attachment_is_explicit(tmp_path):
    loader = _loader(tmp_path)

    parsed = parse_attachment(loader, "attachments/not-present.xlsx")

    assert parsed.readable is False
    assert parsed.missing is True
    assert parsed.error_code == "missing_attachment"
