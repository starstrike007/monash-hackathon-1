from __future__ import annotations

from pathlib import Path

from app.adapters.dataset_loader import DatasetLoader
from app.adapters.document_parsers import parse_attachment
from app.api.schemas.common import CanonicalField, ExtractionState
from app.pipeline.extract import extract_document

REPO_ROOT = Path(__file__).resolve().parents[2]


def _field(extraction, name: CanonicalField):
    return next(field for field in extraction.fields if field.field_name == name)


def test_txt_location_is_line_number(fixture_loader: DatasetLoader) -> None:
    parsed = parse_attachment(fixture_loader, "attachments/fixture_ok_si.txt")
    extraction = extract_document(parsed)
    shipper = _field(extraction, CanonicalField.SHIPPER)
    assert shipper.state == ExtractionState.FOUND
    assert shipper.evidence.line == 2
    assert shipper.evidence.quoted_text == "Shipper: Meridian Pulp Sdn. Bhd."
    assert shipper.evidence.page is None
    assert shipper.evidence.sheet is None


def test_pdf_location_has_page_and_bbox() -> None:
    loader = DatasetLoader(REPO_ROOT / "data")
    parsed = parse_attachment(loader, "attachments/email_059_SI.pdf")
    extraction = extract_document(parsed)
    shipper = _field(extraction, CanonicalField.SHIPPER)
    assert shipper.state == ExtractionState.FOUND
    assert shipper.evidence.page == 1
    assert shipper.evidence.bbox is not None
    for key in ("x0", "x1", "top", "bottom"):
        assert 0.0 <= shipper.evidence.bbox[key] <= 1.0


def test_xlsx_location_has_sheet_and_cell(tmp_path: Path) -> None:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "S.I."
    sheet["A4"] = "Shipper"
    sheet["B4"] = "Test Exporter Pte Ltd"
    (tmp_path / "attachments").mkdir()
    workbook.save(tmp_path / "attachments" / "sample.xlsx")

    loader = DatasetLoader(tmp_path)
    parsed = parse_attachment(loader, "attachments/sample.xlsx")
    extraction = extract_document(parsed)
    shipper = _field(extraction, CanonicalField.SHIPPER)
    assert shipper.state == ExtractionState.FOUND
    assert shipper.evidence.sheet == "S.I."
    assert shipper.evidence.cell == "B4"


def test_docx_location_has_table_and_row_index(tmp_path: Path) -> None:
    from docx import Document

    document = Document()
    table = document.add_table(rows=0, cols=2)
    row = table.add_row()
    row.cells[0].text = "Shipper"
    row.cells[1].text = "Test Exporter Pte Ltd"
    (tmp_path / "attachments").mkdir()
    document.save(tmp_path / "attachments" / "sample.docx")

    loader = DatasetLoader(tmp_path)
    parsed = parse_attachment(loader, "attachments/sample.docx")
    extraction = extract_document(parsed)
    shipper = _field(extraction, CanonicalField.SHIPPER)
    assert shipper.state == ExtractionState.FOUND
    assert shipper.evidence.table_index == 0
    assert shipper.evidence.row_index == 0
